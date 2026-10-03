"""Staff-only administrative views."""

from __future__ import annotations

import logging
from io import BytesIO

import openpyxl
from django.conf import settings
from django.contrib.auth import logout
from django.contrib.auth.decorators import user_passes_test
from django.contrib.auth.models import User
from django.core.exceptions import ValidationError
from django.db import transaction
from django.db.models import Avg, Count, Max, Prefetch
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404
from django.utils import timezone
from django.views.decorators.http import require_http_methods

from common.i18n import translate as t

from apps.accounts.models import UserProfile
from ..forms import (
    AdminAnswerForm,
    AdminBookForm,
    AdminCommentReplyForm,
    AdminPasswordForm,
    AdminResponseForm,
    AdminUserUpdateForm,
    CategoryMutationForm,
    NBOKCategoryCreateForm,
    PhysicalApprovalForm,
    SendBookLinkForm,
)
from ..services.request_data import form_errors
from ..models import (
    Book, BookComment, BookRequest, Category, Message, PhysicalBookBorrow,
    PhysicalBookRequest, Question, UserNotification, BackgroundJob,
)
from ..services import catalog as catalog_service
from ..services import nbok as nbok_service
from ..services.email import record_email as _record_email
from ..services.email import send_book_email, send_library_email as _send_library_email, site_url as _site_url
from ..services.notifications import mark_admin_read as _mark_admin_read
from ..services.notifications import notify as _notify
from ..services.notifications import respond as _respond
from ..services.imports import normalize_import_rows, rows_from_xlsx
from ..services.jobs import enqueue_job, job_public_result
from ..services.request_data import parse_json_array, parse_json_object
from ..services.serialization import serialize_book as _serialize_book
from ..storage import delete_library_asset, save_library_asset

logger = logging.getLogger(__name__)


DASHBOARD_ADMIN_GROUP = "Library Dashboard Admin"

def is_staff_user(user):
    return (
        user.is_authenticated
        and (
            user.is_superuser
            or user.is_staff
            or user.groups.filter(name=DASHBOARD_ADMIN_GROUP).exists()
        )
    )


def _json_body(request):
    return parse_json_object(request)


def _validated_json_form(request, form_class, **kwargs):
    form = form_class(_json_body(request), **kwargs)
    if not form.is_valid():
        raise ValidationError(form_errors(form))
    return form.cleaned_data


@user_passes_test(is_staff_user, login_url="/accounts/login/")
@require_http_methods(["POST"])
def api_upload(request):
    files = request.FILES.getlist("files")
    requested_type = request.POST.get("fileType", "")
    if requested_type not in {"image", "cover", "pdf"}:
        return JsonResponse({"success": False, "message": t("admin.upload.invalid_type")}, status=400)
    asset_type = "cover" if requested_type in {"image", "cover"} else "pdf"
    if not files:
        return JsonResponse({"success": False, "message": t("admin.upload.no_files")}, status=400)
    if len(files) > 20:
        return JsonResponse({"success": False, "message": t("admin.upload.too_many")}, status=400)
    links = []
    try:
        for uploaded in files:
            url = save_library_asset(uploaded, asset_type)
            links.append({"fileName": uploaded.name, "link": url})
    except ValidationError as exc:
        for item in links:
            _cleanup_asset(item["link"], asset_type)
        return JsonResponse({"success": False, "message": "; ".join(exc.messages)}, status=400)
    except Exception:
        for item in links:
            _cleanup_asset(item["link"], asset_type)
        logger.exception("library asset storage failed")
        return JsonResponse({"success": False, "message": t("admin.upload.save_failed")}, status=500)
    return JsonResponse({"success": True, "links": links})


def _admin_unread_counts():
    messages = Message.objects.filter(is_read=False).count()
    digital = BookRequest.objects.filter(admin_read_at__isnull=True).count()
    physical = PhysicalBookRequest.objects.filter(admin_read_at__isnull=True).count()
    return {'messages': messages, 'digital': digital, 'physical': physical,
            'total': messages + digital + physical}


@user_passes_test(is_staff_user, login_url='/accounts/login/')
@require_http_methods(['GET'])
def api_admin_dashboard(request):
    activities = []
    for item in Message.objects.order_by('-created_at')[:4]:
        activities.append({'type': 'message', 'name': item.name, 'content': item.message[:100],
                           'created_at': item.created_at.isoformat()})
    for item in BookRequest.objects.select_related('book').order_by('-created_at')[:4]:
        activities.append({'type': 'book_request',
                           'book_title': item.book.title if item.book else t("admin.dashboard.new_request"),
                           'created_at': item.created_at.isoformat()})
    for item in PhysicalBookRequest.objects.select_related('book').order_by('-created_at')[:4]:
        activities.append({'type': 'physical_request', 'name': item.name,
                           'book_title': item.book.title, 'created_at': item.created_at.isoformat()})
    for item in BookComment.objects.select_related('book').order_by('-created_at')[:4]:
        activities.append({'type': 'comment', 'name': item.name,
                           'book_title': item.book.title, 'created_at': item.created_at.isoformat()})
    activities.sort(key=lambda row: row['created_at'], reverse=True)
    return JsonResponse({
        'success': True, 'totalBooks': Book.objects.count(),
        'totalRequests': BookRequest.objects.filter(status='pending').count(),
        'totalPhysicalRequests': PhysicalBookRequest.objects.filter(status='pending').count(),
        'totalBorrowed': PhysicalBookBorrow.objects.filter(is_returned=False).count(),
        'unreadTotal': _admin_unread_counts()['total'], 'activities': activities[:12],
    })


def _all_books_data():
    parent_map = dict(Category.objects.values_list('id', 'parent_id'))
    queryset = Book.objects.select_related('category', 'nbok_category').prefetch_related(
        Prefetch('borrows', queryset=PhysicalBookBorrow.objects.filter(is_returned=False), to_attr='active_borrows_cache')
    ).annotate(
        average_rating=Avg('ratings__rating'), rating_total=Count('ratings', distinct=True),
    ).order_by('-created_at')
    return [_serialize_book(book, parent_map, include_admin=True) for book in queryset]


@user_passes_test(is_staff_user, login_url='/accounts/login/')
@require_http_methods(['GET'])
def api_admin_books(request):
    return JsonResponse({'success': True, 'data': _all_books_data()})


@user_passes_test(is_staff_user, login_url='/accounts/login/')
@require_http_methods(['GET'])
def api_admin_book_detail(request, book_id):
    book = get_object_or_404(Book.objects.select_related('category', 'nbok_category'), id=book_id)
    parent_map = dict(Category.objects.values_list('id', 'parent_id'))
    return JsonResponse({'success': True, 'data': _serialize_book(book, parent_map, include_admin=True)})


def _category_from_id(value):
    if value in (None, ""):
        return None
    if not str(value).isdigit():
        raise ValidationError(t("admin.category.not_found"))
    category = Category.objects.filter(id=int(value)).first()
    if category is None:
        raise ValidationError(t("admin.category.not_found"))
    return category


def _cleanup_asset(url: str, asset_type: str) -> None:
    if not url:
        return
    try:
        delete_library_asset(url, asset_type)
    except Exception:
        logger.warning(
            "managed asset cleanup failed",
            exc_info=True,
            extra={"asset_type": asset_type},
        )


def _save_book_uploads(data) -> tuple[dict[str, str], list[tuple[str, str]]]:
    uploaded: dict[str, str] = {}
    created: list[tuple[str, str]] = []
    try:
        if data.get("cover"):
            uploaded["cover_url"] = save_library_asset(data["cover"], "cover")
            created.append((uploaded["cover_url"], "cover"))
        if data.get("book_pdf"):
            uploaded["pdf_url"] = save_library_asset(data["book_pdf"], "pdf")
            created.append((uploaded["pdf_url"], "pdf"))
    except Exception:
        for url, asset_type in created:
            _cleanup_asset(url, asset_type)
        raise
    return uploaded, created


def _book_form_values(request, book=None):
    form = AdminBookForm(request.POST, request.FILES)
    if not form.is_valid():
        raise ValidationError(form_errors(form))
    data = form.cleaned_data
    uploads, created = _save_book_uploads(data)
    try:
        book_type = data["book_type"]
        count = 0 if book_type == "pdf" else (data.get("physical_count") or 0)
        cover_url = uploads.get("cover_url") or (data.get("cover_url") or "").strip() or (book.cover_url if book else "")
        pdf_url = uploads.get("pdf_url") or (book.pdf_url if book else "")
        if book_type == "physical":
            pdf_url = ""
        if book_type in {"pdf", "both"} and not pdf_url:
            raise ValidationError(t("admin.book.pdf_upload_required"))
        if "category_id" in request.POST:
            legacy_category = _category_from_id(data.get("category_id"))
        else:
            legacy_category = book.category if book else None
        if "nbok_category_id" in request.POST or "nbok_category_level" in request.POST:
            nbok_category, nbok_category_level = nbok_service.resolve_selection(
                data.get("nbok_category_id"), data.get("nbok_category_level"),
            )
        else:
            nbok_category = book.nbok_category if book else None
            nbok_category_level = book.nbok_category_level if book else None
        values = {
            "title": data["title"].strip(),
            "author": data["author"].strip(),
            "translator": (data.get("translator") or "").strip() or None,
            "publisher": (data.get("publisher") or "").strip() or None,
            "library": (data.get("library") or "").strip() or None,
            "category": legacy_category,
            "nbok_category": nbok_category,
            "nbok_category_level": nbok_category_level,
            "summary": (data.get("summary") or "").strip() or None,
            "cover_url": cover_url or None,
            "pdf_url": pdf_url or None,
            "physical_count": count,
            "physical_location": (data.get("physical_location") or "").strip() or None,
            "tags": data.get("tags") or [],
        }
        return values, created
    except Exception:
        for url, asset_type in created:
            _cleanup_asset(url, asset_type)
        raise


def _remove_replaced_book_assets(old_cover: str, old_pdf: str, book) -> None:
    if old_cover and old_cover != (book.cover_url or ""):
        _cleanup_asset(old_cover, "cover")
    if old_pdf and old_pdf != (book.pdf_url or ""):
        _cleanup_asset(old_pdf, "pdf")


@user_passes_test(is_staff_user, login_url='/accounts/login/')
@require_http_methods(['POST'])
def api_admin_book_add(request):
    created_assets = []
    try:
        values, created_assets = _book_form_values(request)
        tags = values.pop("tags")
        with transaction.atomic():
            book = Book.objects.create(
                physical_available=values["physical_count"],
                tags=tags,
                **values,
            )
        return JsonResponse({
            "success": True,
            "message": t("admin.book.created", title=book.title),
            "book_id": book.id,
        })
    except ValidationError as exc:
        return JsonResponse({"success": False, "message": "; ".join(exc.messages)}, status=400)
    except Exception:
        for url, asset_type in created_assets:
            _cleanup_asset(url, asset_type)
        logger.exception("book create failed")
        return JsonResponse({"success": False, "message": t("admin.book.create_failed")}, status=500)


@user_passes_test(is_staff_user, login_url='/accounts/login/')
@require_http_methods(['POST'])
def api_admin_book_update(request, book_id):
    existing = get_object_or_404(Book, id=book_id)
    created_assets = []
    old_cover, old_pdf = existing.cover_url or "", existing.pdf_url or ""
    try:
        with transaction.atomic():
            book = Book.objects.select_for_update().get(id=existing.id)
            values, created_assets = _book_form_values(request, book)
            tags = values.pop("tags")
            active_borrows = book.borrows.filter(is_returned=False).count()
            if values["physical_count"] < active_borrows:
                raise ValidationError(t("admin.book.active_borrow_count_guard", count=active_borrows))
            for field, value in values.items():
                setattr(book, field, value)
            book.physical_available = max(0, book.physical_count - active_borrows)
            book.tags = tags
            book.save()
        _remove_replaced_book_assets(old_cover, old_pdf, book)
        return JsonResponse({
            "success": True,
            "message": t("admin.book.updated", title=book.title),
        })
    except ValidationError as exc:
        for url, asset_type in created_assets:
            _cleanup_asset(url, asset_type)
        return JsonResponse({"success": False, "message": "; ".join(exc.messages)}, status=400)
    except Exception:
        for url, asset_type in created_assets:
            _cleanup_asset(url, asset_type)
        logger.exception("book update failed", extra={"book_id": book_id})
        return JsonResponse({"success": False, "message": t("admin.book.update_failed")}, status=500)


@user_passes_test(is_staff_user, login_url='/accounts/login/')
@require_http_methods(['POST'])
def api_admin_book_delete(request, book_id):
    existing = get_object_or_404(Book, id=book_id)
    cover_url, pdf_url = existing.cover_url or "", existing.pdf_url or ""
    try:
        with transaction.atomic():
            book = Book.objects.select_for_update().get(id=existing.id)
            if book.borrows.filter(is_returned=False).exists():
                return JsonResponse({
                    "success": False,
                    "message": t("admin.book.active_borrow_delete_guard"),
                }, status=400)
            book.delete()
        _cleanup_asset(cover_url, "cover")
        _cleanup_asset(pdf_url, "pdf")
        return JsonResponse({"success": True, "message": t("admin.book.deleted")})
    except Exception:
        logger.exception("book delete failed", extra={"book_id": book_id})
        return JsonResponse({"success": False, "message": t("admin.book.delete_failed")}, status=500)


PHYSICAL_IMPORT_HEADERS = [
    t("books.book.field.title"), t("common.labels.author"), t("common.labels.translator"),
    t("common.labels.publisher"), t("books.book.field.library"), t("books.category.model.single"),
    t("books.book.field.summary"), t("books.book.field.tags"), t("books.book.field.physical_count"),
    t("books.book.field.physical_location"), t("admin.import.header.cover_image"),
]



@user_passes_test(is_staff_user, login_url='/accounts/login/')
@require_http_methods(['POST'])
def api_admin_books_import(request):
    excel = request.FILES.get('excel')
    try:
        if excel:
            rows = rows_from_xlsx(excel)
        else:
            rows = normalize_import_rows(parse_json_array(request))
    except ValidationError as exc:
        return JsonResponse({'success': False, 'message': '; '.join(exc.messages)}, status=400)

    job = enqueue_job(
        BackgroundJob.TYPE_PHYSICAL_IMPORT,
        {'rows': rows},
        created_by=request.user,
        max_attempts=3,
    )
    data = job_public_result(job)
    if job.status == BackgroundJob.STATUS_COMPLETED:
        # Preserve the legacy synchronous response contract in eager/test mode.
        data.update(job.result)
    data.update({
        'success': True,
        'message': t("admin.import.queued") if job.status != BackgroundJob.STATUS_COMPLETED else _import_result_message(job.result),
    })
    return JsonResponse(data, status=202 if job.status != BackgroundJob.STATUS_COMPLETED else 200)


def _import_result_message(result):
    success_count = int(result.get('success_count', 0))
    error_count = int(result.get('error_count', 0))
    text = t("admin.import.success_count", count=success_count)
    return text + (t("admin.import.error_count_suffix", count=error_count) if error_count else '')


@user_passes_test(is_staff_user, login_url='/accounts/login/')
@require_http_methods(['GET'])
def api_admin_import_job(request, job_id):
    job = get_object_or_404(
        BackgroundJob,
        id=job_id,
        job_type=BackgroundJob.TYPE_PHYSICAL_IMPORT,
        created_by=request.user,
    )
    data = job_public_result(job)
    data['success'] = job.status != BackgroundJob.STATUS_FAILED
    if job.status == BackgroundJob.STATUS_COMPLETED:
        data['message'] = _import_result_message(job.result)
    elif job.status == BackgroundJob.STATUS_FAILED:
        data['message'] = t("admin.import.failed_after_retries")
    else:
        data['message'] = t("admin.import.processing")
    return JsonResponse(data)


@user_passes_test(is_staff_user, login_url='/accounts/login/')
@require_http_methods(['GET'])
def api_admin_physical_import_template(request):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = t("admin.import.sheet_title")
    sheet.append(PHYSICAL_IMPORT_HEADERS)
    sheet.append([t("admin.import.sample.title"), t("admin.import.sample.author"), t("admin.import.sample.translator"),
                  t("admin.import.sample.publisher"), t("admin.import.sample.library"), t("admin.import.sample.category"),
                  t("admin.import.sample.summary"), t("admin.import.sample.tags"), 2, t("admin.import.sample.location"),
                  'https://cdn.example.com/library/Cover/example.jpg'])
    for index, width in enumerate([28, 24, 20, 20, 18, 24, 42, 28, 18, 20, 52], start=1):
        sheet.column_dimensions[openpyxl.utils.get_column_letter(index)].width = width
    for cell in sheet[1]:
        cell.font = openpyxl.styles.Font(bold=True, color='FFFFFF')
        cell.fill = openpyxl.styles.PatternFill('solid', fgColor='0F766E')
    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(output.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="library-physical-books-template.xlsx"'
    return response


@user_passes_test(is_staff_user, login_url='/accounts/login/')
@require_http_methods(['GET'])
def api_admin_physical_requests(request):
    queryset = PhysicalBookRequest.objects.select_related('book').order_by('-created_at')[:300]
    data = []
    for item in queryset:
        data.append({
            'id': item.id, 'book_id': item.book_id, 'book_title': item.book.title,
            'book_author': item.book.author, 'book_cover': item.book.cover_url,
            'name': item.name, 'email': item.email, 'phone': item.phone,
            'message': item.message, 'status': item.status,
            'is_read': bool(item.admin_read_at), 'admin_response': item.admin_response,
            'created_at': item.created_at.isoformat(),
        })
    return JsonResponse({'success': True, 'data': data})


@user_passes_test(is_staff_user, login_url='/accounts/login/')
@require_http_methods(['POST'])
def api_admin_physical_requests_mark_read(request):
    queryset = PhysicalBookRequest.objects.select_related('book').filter(admin_read_at__isnull=True).order_by('-created_at')[:300]
    for item in queryset:
        _mark_admin_read(item, t("admin.borrow.read_label", title=item.book.title))
    return JsonResponse({'success': True, 'marked': len(queryset)})


@user_passes_test(is_staff_user, login_url='/accounts/login/')
@require_http_methods(['POST'])
def api_admin_approve_physical_request(request):
    try:
        data = _validated_json_form(request, PhysicalApprovalForm)
    except ValidationError as exc:
        return JsonResponse({'success': False, 'message': '; '.join(exc.messages)}, status=400)
    with transaction.atomic():
        item = get_object_or_404(
            PhysicalBookRequest.objects.select_for_update().select_related('book'),
            id=data['request_id'],
        )
        book = Book.objects.select_for_update().get(id=item.book_id)
        if item.status != 'pending':
            return JsonResponse({'success': False, 'message': t("admin.borrow.already_reviewed")}, status=400)
        if book.physical_available <= 0:
            return JsonResponse({'success': False, 'message': t("admin.borrow.unavailable")}, status=400)
        PhysicalBookBorrow.objects.create(
            book=book, borrower_name=item.name, borrower_email=item.email,
            borrower_phone=item.phone, borrow_date=data['borrow_date'], return_date=data['return_date'],
            notes=(data.get('notes') or '').strip(),
        )
        item.status = 'approved'
        item.save(update_fields=['status'])
        book.physical_available -= 1
        book.save(update_fields=['physical_available'])
        _respond(
            item, 'approved', t("admin.borrow.approved_title"),
            t("admin.borrow.approved_message", title=book.title),
        )
        email_sent = _send_library_email(
            _record_email(item),
            t("admin.borrow.email_subject", title=book.title),
            t("admin.borrow.email_heading"),
            t("admin.borrow.email_message"),
            eyebrow=t("admin.borrow.email_eyebrow"),
            action_label=t("admin.borrow.email_action"),
            action_url=_site_url('/accounts/profile/'),
            details=[{'label': t("common.labels.book"), 'value': book.title}],
        )
    message = t("admin.borrow.completed_with_email") if email_sent else t("admin.borrow.completed_without_email")
    return JsonResponse({'success': True, 'message': message, 'email_sent': email_sent})


@user_passes_test(is_staff_user, login_url='/accounts/login/')
@require_http_methods(['POST'])
def api_admin_reject_physical_request(request, request_id):
    item = get_object_or_404(PhysicalBookRequest.objects.select_related('book'), id=request_id)
    if item.status != 'pending':
        return JsonResponse({'success': False, 'message': t("admin.borrow.already_reviewed")}, status=409)
    item.status = 'rejected'
    item.save(update_fields=['status'])
    try:
        payload = _json_body(request)
    except ValidationError as exc:
        return JsonResponse({'success': False, 'message': '; '.join(exc.messages)}, status=400)
    response_text = (
        payload.get('response')
        or t("admin.borrow.rejected_default", title=item.book.title)
    ).strip()
    _respond(item, 'rejected', t("admin.borrow.reviewed_title"), response_text)
    email_sent = _send_library_email(
        _record_email(item),
        t("admin.borrow.result_subject", title=item.book.title),
        t("admin.borrow.result_heading"),
        response_text,
        eyebrow=t("admin.borrow.result_eyebrow"),
        action_label=t("admin.requests.action"),
        action_url=_site_url('/accounts/profile/'),
    )
    message = t("admin.delivery.result_with_email") if email_sent else t("admin.delivery.result_without_email")
    return JsonResponse({'success': True, 'message': message, 'email_sent': email_sent})


@user_passes_test(is_staff_user, login_url='/accounts/login/')
@require_http_methods(['GET'])
def api_admin_borrowed_books(request):
    data, today = [], timezone.now().date()
    for item in PhysicalBookBorrow.objects.select_related('book').filter(is_returned=False).order_by('-borrow_date'):
        data.append({
            'id': item.id, 'book_id': item.book_id, 'book_title': item.book.title,
            'book_author': item.book.author, 'book_cover': item.book.cover_url,
            'borrower_name': item.borrower_name, 'borrower_email': item.borrower_email,
            'borrower_phone': item.borrower_phone, 'borrow_date': item.borrow_date.isoformat(),
            'return_date': item.return_date.isoformat(), 'notes': item.notes,
            'is_overdue': item.return_date < today,
            'days_overdue': max(0, (today - item.return_date).days),
        })
    return JsonResponse({'success': True, 'data': data})


@user_passes_test(is_staff_user, login_url='/accounts/login/')
@require_http_methods(['GET'])
def api_admin_borrow_history(request, book_id):
    data = [{
        'id': item.id, 'borrower_name': item.borrower_name,
        'borrower_email': item.borrower_email, 'borrower_phone': item.borrower_phone,
        'borrow_date': item.borrow_date.isoformat(), 'return_date': item.return_date.isoformat(),
        'actual_return_date': item.actual_return_date.isoformat() if item.actual_return_date else None,
        'is_returned': item.is_returned, 'notes': item.notes,
    } for item in PhysicalBookBorrow.objects.filter(book_id=book_id).order_by('-borrow_date')]
    return JsonResponse({'success': True, 'data': data})


@user_passes_test(is_staff_user, login_url='/accounts/login/')
@require_http_methods(['POST'])
def api_admin_return_book(request, borrow_id):
    with transaction.atomic():
        borrow = get_object_or_404(PhysicalBookBorrow.objects.select_for_update(), id=borrow_id)
        if borrow.is_returned:
            return JsonResponse({'success': False, 'message': t("admin.return.already_recorded")}, status=400)
        borrow.is_returned = True
        borrow.actual_return_date = timezone.now().date()
        borrow.save(update_fields=['is_returned', 'actual_return_date'])
        book = Book.objects.select_for_update().get(id=borrow.book_id)
        book.physical_available = min(book.physical_count, book.physical_available + 1)
        book.save(update_fields=['physical_available'])
        email_sent = _send_library_email(
            borrow.borrower_email,
            t("admin.return.email_subject", title=book.title),
            t("admin.return.email_heading"),
            t("admin.return.email_message"),
            eyebrow=t("admin.return.email_eyebrow"),
            action_label=t("admin.return.email_action"),
            action_url=_site_url('/library/books/'),
            details=[{'label': t("common.labels.book"), 'value': book.title}],
        )
    message = t("admin.return.completed_with_email") if email_sent else t("admin.return.completed_without_email")
    return JsonResponse({'success': True, 'message': message, 'email_sent': email_sent})


@user_passes_test(is_staff_user, login_url='/accounts/login/')
@require_http_methods(['POST'])
def api_admin_borrow_history_delete(request, borrow_id):
    borrow = get_object_or_404(PhysicalBookBorrow, id=borrow_id)
    if not borrow.is_returned:
        return JsonResponse({'success': False, 'message': t("admin.return.active_required")}, status=400)
    borrow.delete()
    return JsonResponse({'success': True, 'message': t("admin.return.history_deleted")})


@user_passes_test(is_staff_user, login_url='/accounts/login/')
@require_http_methods(['GET'])
def api_admin_book_requests(request):
    queryset = BookRequest.objects.select_related('book').order_by('-created_at')[:300]
    data = []
    for item in queryset:
        title = item.book.title if item.book else t("books.digital_request.new_book")
        data.append({
            'id': item.id, 'email': item.email, 'name': item.name,
            'message': item.message, 'status': item.status,
            'is_read': bool(item.admin_read_at), 'admin_response': item.admin_response,
            'created_at': item.created_at.isoformat(), 'book_id': item.book_id,
            'book_title': title, 'book_author': item.book.author if item.book else '',
            'book_cover': item.book.cover_url if item.book else '',
        })
    return JsonResponse({'success': True, 'data': data})


@user_passes_test(is_staff_user, login_url='/accounts/login/')
@require_http_methods(['POST'])
def api_admin_book_requests_mark_read(request):
    queryset = BookRequest.objects.select_related('book').filter(admin_read_at__isnull=True).order_by('-created_at')[:300]
    for item in queryset:
        title = item.book.title if item.book else t("books.digital_request.new_book")
        _mark_admin_read(item, t("admin.digital.read_label", title=title))
    return JsonResponse({'success': True, 'marked': len(queryset)})


@user_passes_test(is_staff_user, login_url='/accounts/login/')
@require_http_methods(['GET'])
def api_admin_book_requests_pending_count(request):
    return JsonResponse({'success': True,
                         'count': BookRequest.objects.filter(admin_read_at__isnull=True).count()})


@user_passes_test(is_staff_user, login_url='/accounts/login/')
@require_http_methods(['POST'])
def api_admin_approve_book_request(request, request_id):
    item = get_object_or_404(BookRequest.objects.select_related('book'), id=request_id)
    if item.status != 'pending':
        return JsonResponse({'success': False, 'message': t("admin.borrow.already_reviewed")}, status=409)
    item.status = 'approved'
    item.save(update_fields=['status'])
    sent = bool(item.book and send_book_email(item.book, item.email))
    title = item.book.title if item.book else t("admin.digital.requested_book")
    response_text = t("admin.digital.approved_response", title=title)
    if sent:
        response_text += t("admin.digital.email_sent_suffix")
    else:
        response_text += t("admin.digital.email_failed_suffix")
    _respond(item, 'approved', t("admin.digital.approved_title"), response_text)
    return JsonResponse({'success': True, 'message': response_text, 'email_sent': sent})


@user_passes_test(is_staff_user, login_url='/accounts/login/')
@require_http_methods(['POST'])
def api_admin_reject_book_request(request, request_id):
    item = get_object_or_404(BookRequest.objects.select_related('book'), id=request_id)
    if item.status != 'pending':
        return JsonResponse({'success': False, 'message': t("admin.borrow.already_reviewed")}, status=409)
    item.status = 'rejected'
    item.save(update_fields=['status'])
    title = item.book.title if item.book else t("admin.digital.requested_book")
    try:
        payload = _json_body(request)
    except ValidationError as exc:
        return JsonResponse({'success': False, 'message': '; '.join(exc.messages)}, status=400)
    response_text = (
        payload.get('response')
        or t("admin.digital.rejected_default", title=title)
    ).strip()
    _respond(item, 'rejected', t("admin.digital.reviewed_title"), response_text)
    email_sent = _send_library_email(
        _record_email(item),
        t("admin.digital.result_subject", title=title),
        t("admin.digital.result_heading"),
        response_text,
        eyebrow=t("admin.digital.result_eyebrow"),
        action_label=t("admin.requests.action"),
        action_url=_site_url('/accounts/profile/'),
    )
    message = t("admin.delivery.result_with_email") if email_sent else t("admin.delivery.result_without_email")
    return JsonResponse({'success': True, 'message': message, 'email_sent': email_sent})


@user_passes_test(is_staff_user, login_url='/accounts/login/')
@require_http_methods(['POST'])
def api_admin_book_request_delete(request, request_id):
    get_object_or_404(BookRequest, id=request_id).delete()
    return JsonResponse({'success': True, 'message': t("admin.digital.deleted")})


@user_passes_test(is_staff_user, login_url='/accounts/login/')
@require_http_methods(['GET'])
def api_admin_messages(request):
    data = [{
        'id': item.id, 'name': item.name, 'email': item.email,
        'message': item.message, 'is_read': item.is_read,
        'admin_response': item.admin_response,
        'responded_at': item.responded_at.isoformat() if item.responded_at else None,
        'created_at': item.created_at.isoformat(),
    } for item in Message.objects.order_by('-created_at')[:300]]
    return JsonResponse({'success': True, 'data': data})


@user_passes_test(is_staff_user, login_url='/accounts/login/')
@require_http_methods(['POST'])
def api_admin_message_read(request, message_id):
    item = get_object_or_404(Message, id=message_id)
    _mark_admin_read(item, t("admin.message.read_label"))
    return JsonResponse({'success': True, 'message': t("admin.message.read")})


@user_passes_test(is_staff_user, login_url='/accounts/login/')
@require_http_methods(['POST'])
def api_admin_message_reply(request, message_id):
    item = get_object_or_404(Message, id=message_id)
    try:
        payload = _json_body(request)
    except ValidationError as exc:
        return JsonResponse({'success': False, 'message': '; '.join(exc.messages)}, status=400)
    response_text = (payload.get('response') or '').strip()
    if not response_text:
        return JsonResponse({'success': False, 'message': t("admin.message.response_required")}, status=400)
    _respond(item, 'answered', t("admin.message.response_title"), response_text)
    email_sent = _send_library_email(
        _record_email(item),
        t("admin.message.email_subject"),
        t("admin.message.email_heading"),
        response_text,
        eyebrow=t("admin.message.email_eyebrow"),
        action_label=t("admin.message.email_action"),
        action_url=_site_url('/accounts/notifications/'),
        details=[{'label': t("admin.message.read_label"), 'value': item.message[:220]}],
    )
    message = t("admin.message.response_with_email") if email_sent else t("admin.message.response_without_email")
    return JsonResponse({'success': True, 'message': message, 'email_sent': email_sent})


@user_passes_test(is_staff_user, login_url='/accounts/login/')
@require_http_methods(['POST'])
def api_admin_message_delete(request, message_id):
    get_object_or_404(Message, id=message_id).delete()
    return JsonResponse({'success': True, 'message': t("admin.message.deleted")})


@user_passes_test(is_staff_user, login_url='/accounts/login/')
@require_http_methods(['GET'])
def api_admin_messages_unread_count(request):
    counts = _admin_unread_counts()
    return JsonResponse({'success': True, 'count': counts['messages'], **counts})


@user_passes_test(is_staff_user, login_url='/accounts/login/')
@require_http_methods(['GET'])
def api_admin_unread_counts(request):
    return JsonResponse({'success': True, **_admin_unread_counts()})


@user_passes_test(is_staff_user, login_url='/accounts/login/')
@require_http_methods(['GET'])
def api_admin_users(request):
    users = User.objects.select_related('profile').annotate(
        digital_count=Count('digital_book_requests', distinct=True),
        physical_count=Count('physical_book_requests', distinct=True),
        digital_last=Max('digital_book_requests__created_at'),
        physical_last=Max('physical_book_requests__created_at'),
    ).order_by('-date_joined')
    data = []
    for user in users:
        last_values = [value for value in [user.digital_last, user.physical_last] if value]
        profile = getattr(user, 'profile', None)
        data.append({
            'id': user.id, 'username': user.username,
            'name': user.get_full_name() or user.username,
            'first_name': user.first_name, 'last_name': user.last_name,
            'email': user.email, 'phone': profile.phone_number if profile else '',
            'is_active': user.is_active, 'is_staff': user.is_staff,
            'is_superuser': user.is_superuser,
            'can_delete': user.pk != request.user.pk and not user.is_superuser,
            'date_joined': user.date_joined.isoformat(),
            'request_count': user.digital_count + user.physical_count,
            'last_request': max(last_values).isoformat() if last_values else None,
        })
    return JsonResponse({'success': True, 'data': data})


@user_passes_test(is_staff_user, login_url='/accounts/login/')
@require_http_methods(['POST'])
def api_admin_user_update(request, user_id):
    user = get_object_or_404(User, id=user_id)
    try:
        data = _validated_json_form(request, AdminUserUpdateForm)
    except ValidationError as exc:
        return JsonResponse({'success': False, 'message': '; '.join(exc.messages)}, status=400)
    username = data['username'].strip()
    email = data['email'].strip().lower()
    if User.objects.filter(username=username).exclude(pk=user.pk).exists():
        return JsonResponse({'success': False, 'message': t("admin.user.username_exists")}, status=400)
    if User.objects.filter(email__iexact=email).exclude(pk=user.pk).exists():
        return JsonResponse({'success': False, 'message': t("admin.user.email_exists")}, status=400)
    user.username = username
    user.email = email
    user.first_name = (data.get('first_name') or '').strip()
    user.last_name = (data.get('last_name') or '').strip()
    if user.pk != request.user.pk:
        if data.get('is_active') is not None:
            user.is_active = data['is_active']
        if request.user.is_superuser and not user.is_superuser and data.get('is_staff') is not None:
            user.is_staff = data['is_staff']
    user.save(update_fields=['username', 'email', 'first_name', 'last_name', 'is_active', 'is_staff'])
    profile, _ = UserProfile.objects.get_or_create(user=user)
    profile.phone_number = (data.get('phone') or '').strip()
    profile.save(update_fields=['phone_number'])
    return JsonResponse({'success': True, 'message': t("admin.user.updated")})


@user_passes_test(is_staff_user, login_url='/accounts/login/')
@require_http_methods(['POST'])
def api_admin_user_password(request, user_id):
    user = get_object_or_404(User, id=user_id)
    try:
        data = _validated_json_form(request, AdminPasswordForm, user=user)
    except ValidationError as exc:
        return JsonResponse({'success': False, 'message': '; '.join(exc.messages)}, status=400)
    user.set_password(data['password'])
    user.save(update_fields=['password'])
    return JsonResponse({'success': True, 'message': t("admin.user.password_set")})


@user_passes_test(is_staff_user, login_url='/accounts/login/')
@require_http_methods(['POST'])
def api_admin_user_delete(request, user_id):
    user = get_object_or_404(User, id=user_id)
    if user.pk == request.user.pk:
        return JsonResponse({'success': False, 'message': t("admin.user.current_admin_delete_guard")}, status=400)
    if user.is_superuser:
        return JsonResponse({'success': False, 'message': t("admin.user.superuser_delete_guard")}, status=400)
    user.delete()
    return JsonResponse({'success': True, 'message': t("admin.user.deleted")})


@user_passes_test(is_staff_user, login_url='/accounts/login/')
@require_http_methods(['GET'])
def api_admin_comments(request):
    comments = BookComment.objects.select_related('book').filter(
        parent__isnull=True,
    ).prefetch_related('replies').order_by('-created_at')
    data = [{
        'id': item.id, 'name': item.name, 'email': item.email,
        'text': item.comment, 'book_id': item.book_id, 'book_title': item.book.title,
        'created_at': item.created_at.isoformat(), 'is_approved': item.is_approved,
        'replies': [
            {'id': reply.id, 'name': reply.name, 'text': reply.comment,
             'date': reply.created_at.isoformat()}
            for reply in item.replies.all()
        ],
    } for item in comments]
    return JsonResponse({'success': True, 'data': data})


@user_passes_test(is_staff_user, login_url='/accounts/login/')
@require_http_methods(['POST'])
def api_admin_comment_reply(request, comment_id):
    parent = get_object_or_404(BookComment.objects.select_related('book'), id=comment_id)
    try:
        data = _validated_json_form(request, AdminCommentReplyForm)
    except ValidationError as exc:
        return JsonResponse({'success': False, 'message': '; '.join(exc.messages)}, status=400)
    response_text = data['text']
    reply = BookComment.objects.create(
        book=parent.book,
        parent=parent,
        name=t("admin.comment.manager_name"),
        email=request.user.email or settings.DEFAULT_FROM_EMAIL,
        comment=response_text,
        is_approved=True,
    )
    _notify(
        parent,
        'answered',
        t("admin.comment.notification_title", title=parent.book.title),
        response_text,
    )
    email_sent = _send_library_email(
        _record_email(parent),
        t("admin.comment.email_subject", title=parent.book.title),
        t("admin.comment.email_heading"),
        response_text,
        eyebrow=t("admin.comment.email_eyebrow"),
        action_label=t("admin.return.email_action"),
        action_url=_site_url('/library/books/'),
        details=[{'label': t("common.labels.book"), 'value': parent.book.title}],
    )
    return JsonResponse({
        'success': True,
        'message': t("admin.comment.response_with_email") if email_sent else t("admin.comment.response_without_email"),
        'reply_id': reply.id,
        'email_sent': email_sent,
    })


@user_passes_test(is_staff_user, login_url='/accounts/login/')
@require_http_methods(['POST'])
def api_admin_comment_delete(request, comment_id):
    get_object_or_404(BookComment, id=comment_id).delete()
    return JsonResponse({'success': True, 'message': t("admin.comment.deleted")})


@user_passes_test(is_staff_user, login_url='/accounts/login/')
@require_http_methods(['GET'])
def api_admin_categories(request):
    return JsonResponse({
        'success': True,
        'data': catalog_service.category_tree(),
    })

@user_passes_test(is_staff_user, login_url='/accounts/login/')
@require_http_methods(['GET'])
def api_admin_categories_book_counts(request):
    data = dict(Book.objects.values_list('category_id').annotate(total=Count('id')))
    return JsonResponse({'success': True, 'data': data})


@user_passes_test(is_staff_user, login_url='/accounts/login/')
@require_http_methods(['GET'])
def api_admin_nbok(request):
    return JsonResponse({'success': True, 'data': nbok_service.tree()})


@user_passes_test(is_staff_user, login_url='/accounts/login/')
@require_http_methods(['POST'])
def api_admin_nbok_create(request):
    try:
        data = _validated_json_form(request, NBOKCategoryCreateForm)
        entry = nbok_service.create_category(
            name=data['name'],
            parent_row_id=data.get('parent_row_id'),
            parent_level=data.get('parent_level'),
        )
    except ValidationError as exc:
        return JsonResponse({'success': False, 'message': '; '.join(exc.messages)}, status=400)
    return JsonResponse({
        'success': True,
        'message': t('admin.nbok.created'),
        'row_id': entry.id,
    })


def _is_descendant(category, candidate):
    current = candidate
    while current:
        if current.id == category.id:
            return True
        current = current.parent
    return False


@user_passes_test(is_staff_user, login_url='/accounts/login/')
@require_http_methods(['POST'])
def api_admin_category_create(request):
    try:
        data = _validated_json_form(request, CategoryMutationForm)
        parent = _category_from_id(data.get('parent_id'))
    except ValidationError as exc:
        return JsonResponse({'success': False, 'message': '; '.join(exc.messages)}, status=400)
    name = data['name'].strip()
    if Category.objects.filter(name=name).exists():
        return JsonResponse({'success': False, 'message': t("admin.category.exists")}, status=400)
    category = Category.objects.create(name=name, parent=parent)
    return JsonResponse({'success': True, 'message': t("admin.category.created"),
                         'category_id': category.id})


@user_passes_test(is_staff_user, login_url='/accounts/login/')
@require_http_methods(['POST'])
def api_admin_category_update(request, category_id):
    category = get_object_or_404(Category, id=category_id)
    try:
        data = _validated_json_form(request, CategoryMutationForm)
        parent = _category_from_id(data.get('parent_id'))
    except ValidationError as exc:
        return JsonResponse({'success': False, 'message': '; '.join(exc.messages)}, status=400)
    name = data['name'].strip()
    if parent and _is_descendant(category, parent):
        return JsonResponse({'success': False, 'message': t("admin.category.parent_invalid")}, status=400)
    if Category.objects.filter(name=name).exclude(id=category.id).exists():
        return JsonResponse({'success': False, 'message': t("admin.category.exists")}, status=400)
    category.name, category.parent = name, parent
    category.save(update_fields=['name', 'parent'])
    return JsonResponse({'success': True, 'message': t("admin.category.updated")})


@user_passes_test(is_staff_user, login_url='/accounts/login/')
@require_http_methods(['POST'])
def api_admin_category_delete(request, category_id):
    category = get_object_or_404(Category, id=category_id)
    if category.children.exists() or category.book_set.exists():
        return JsonResponse({
            'success': False,
            'message': t("admin.category.not_empty"),
        }, status=400)
    category.delete()
    return JsonResponse({'success': True, 'message': t("admin.category.deleted")})


@user_passes_test(is_staff_user, login_url='/accounts/login/')
@require_http_methods(['POST'])
def api_admin_send_book_link(request):
    try:
        data = _validated_json_form(request, SendBookLinkForm)
    except ValidationError as exc:
        return JsonResponse({'success': False, 'message': '; '.join(exc.messages)}, status=400)
    book = get_object_or_404(Book, id=data['book_id'])
    email = data['email'].strip().lower()
    if not book.has_pdf:
        return JsonResponse({'success': False, 'message': t("admin.pdf.missing")}, status=400)
    if not send_book_email(book, email):
        return JsonResponse({
            'success': False,
            'message': t("admin.email.delivery_failed"),
        }, status=500)
    return JsonResponse({'success': True, 'message': t("admin.pdf.link_sent")})


# Legacy compatibility APIs retained for old data; the current admin UI does not expose them.
@user_passes_test(is_staff_user, login_url='/accounts/login/')
@require_http_methods(['GET'])
def api_admin_questions(request):
    data = list(Question.objects.filter(is_approved=False).values(
        'id', 'question', 'book__title', 'created_at',
    ))
    for item in data:
        item['book_title'] = item.pop('book__title')
    return JsonResponse({'success': True, 'data': data})


@user_passes_test(is_staff_user, login_url='/accounts/login/')
@require_http_methods(['POST'])
def api_admin_question_answer(request, question_id):
    question = get_object_or_404(Question, id=question_id)
    try:
        data = _validated_json_form(request, AdminAnswerForm)
    except ValidationError as exc:
        return JsonResponse({'success': False, 'message': '; '.join(exc.messages)}, status=400)
    answer = data['answer'].strip()
    question.answer, question.is_approved, question.answered_at = answer, True, timezone.now()
    question.save(update_fields=['answer', 'is_approved', 'answered_at'])
    return JsonResponse({'success': True, 'message': t("admin.question.answered")})


@user_passes_test(is_staff_user, login_url='/accounts/login/')
@require_http_methods(['POST'])
def api_admin_question_delete(request, question_id):
    get_object_or_404(Question, id=question_id).delete()
    return JsonResponse({'success': True, 'message': t("admin.question.deleted")})


@require_http_methods(['POST'])
def api_logout(request):
    logout(request)
    return JsonResponse({'success': True, 'redirect': '/accounts/login/'})
